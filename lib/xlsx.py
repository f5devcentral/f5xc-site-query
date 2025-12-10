import itertools
import sys
from logging import Logger

from enlighten import get_manager
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill

import lib.const as c

# CONSTS
GREY_FILL = PatternFill(start_color='778491', end_color='778491', fill_type='solid')
GREY_FILL_SECTION = PatternFill(start_color='99AABB', end_color='99AABB', fill_type='solid')
HEADER_FONT = Font(size=18, bold=True, italic=False, color="00000000")
SECTION_FONT = Font(size=14, bold=True, italic=False, color="00000001")
LIGHT_GREY_FILL = PatternFill(start_color='FFD9D9D9', end_color='FFD9D9D9', fill_type='solid')
RIGHT_ALIGNMENT = Alignment(horizontal='right')
LEFT_ALIGNMENT = Alignment(horizontal='left')
WRAP_TEXT_RIGHT_ALIGNMENT = Alignment(horizontal='right', wrap_text=True)
LEFT_CENTER_ALIGNMENT = Alignment(horizontal='left', vertical='center')
LEFT_CENTER_WRAP_TEXT_ALIGNMENT = Alignment(horizontal='left', vertical='center', wrap_text=True)
HEADLINE_CELL_START = 'A1'
HEADLINE_CELLS = f'{HEADLINE_CELL_START}:F1'
COLUMN_DIMENSIONS_A_WIDTH = 30
COLUMN_DIMENSIONS_B_WIDTH = 25


def join_dict_items(data_dict: dict, separator="\n"):
    """
    Joins all key-value pairs in a dictionary into a single string,
    separated by a specified separator (default is a newline).
    """

    formatted_items = separator.join(f"{key}: {value}" for key, value in data_dict.items())

    return formatted_items


class Xlsx(object):
    """
    """

    def __init__(self, site: str = None, file: str = None, logger: Logger = None):
        """

        Parameters
        ----------
        site: str
        file: str
        logger: Logger
        """
        self.logger = logger
        self.file = file
        self.wb = Workbook()
        self.ws = self.wb.active
        self.site = site
        self.must_break = False

    def write(self):
        """

        Returns
        -------

        """
        self.logger.info(f"Writing xlsx file: {self.file}")
        self.wb.save(self.file)
        self.logger.info(f"Writing xlsx file: {self.file}. Done.")

    def build_inventory_summary(self, order: int = None, data: dict = None, title_prefix: str = None):
        """

        Parameters
        ----------
        order: int
        data: dict
        title_prefix: str

        Returns
        -------

        """

        sites = data[c.SITES_KEY]

        # WS Summary Tab
        ws_summary = self.wb.create_sheet("Summary", order)
        ws_summary.column_dimensions['A'].width = COLUMN_DIMENSIONS_A_WIDTH
        ws_summary.column_dimensions['B'].width = COLUMN_DIMENSIONS_B_WIDTH

        def process():
            pbar.update()

            table_data_infrastructure = [
                ("Kind", sites[site]["kind"]),
                ("Provider Type", sites[site]["metadata"]["labels"]["ves.io/provider"] if "ves.io/provider" in sites[site]["metadata"]["labels"] else "Unknown"),
                ("Main Node Count", sites[site]["main_node_count"]),
                ("Worker Node Count", sites[site]["worker_node_count"] if "worker_node_count" in sites[site] else 0),
                ("Node0 CPU Count", sites[site]["nodes"]["node0"]["hw_info"]["cpu"]["cpus"] if "hw_info" in sites[site]["nodes"]["node0"] else 0),
                ("Node0 Memory Size (GB)", round(sites[site]["nodes"]["node0"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in sites[site]["nodes"]["node0"] else 0),
                ("Node0 Interface Count", len(sites[site]["nodes"]["node0"]["interfaces"]) if "interfaces" in sites[site]["nodes"]["node0"] else 0),
            ]

            if "hw_info" in sites[site]["nodes"]["node0"]:
                for storage in sites[site]["nodes"]["node0"]["hw_info"]["storage"]:
                    node0_storage_size = storage["size_gb"]
                    table_data_infrastructure.append((f"Node0 Storage {storage["name"]} Size (GB)", node0_storage_size))

            if sites[site]["main_node_count"] > 1:
                table_data_infrastructure.extend(
                    [
                        ("Node1 CPU Count", sites[site]["nodes"]["node1"]["hw_info"]["cpu"]["cpus"] if "hw_info" in sites[site]["nodes"]["node1"] else 0),
                        ("Node1 Memory Size (GB)", round(sites[site]["nodes"]["node1"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in sites[site]["nodes"]["node1"] else 0),
                        ("Node1 Interface Count", len(sites[site]["nodes"]["node1"]["interfaces"]) if "interfaces" in sites[site]["nodes"]["node1"] else 0),
                    ]
                )

                if "hw_info" in sites[site]["nodes"]["node1"]:
                    for storage in sites[site]["nodes"]["node1"]["hw_info"]["storage"]:
                        node1_storage_size = storage["size_gb"]
                        table_data_infrastructure.append((f"Node1 Storage {storage["name"]} Size (GB)", node1_storage_size))

                table_data_infrastructure.extend(
                    [
                        ("Node2 CPU Count", sites[site]["nodes"]["node2"]["hw_info"]["cpu"]["cpus"] if "hw_info" in sites[site]["nodes"]["node2"] else 0),
                        ("Node2 Memory Size (GB)", round(sites[site]["nodes"]["node2"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in sites[site]["nodes"]["node2"] else 0),
                        ("Node2 Interface Count", len(sites[site]["nodes"]["node2"]["interfaces"]) if "interfaces" in sites[site]["nodes"]["node2"] else 0),
                    ]
                )

                if "hw_info" in sites[site]["nodes"]["node2"]:
                    for storage in sites[site]["nodes"]["node2"]["hw_info"]["storage"]:
                        node2_storage_size = storage["size_gb"]
                        table_data_infrastructure.append((f"Node2 Storage {storage["name"]} Size (GB)", node2_storage_size))

            lbs = 0
            ops = 0
            proxies = 0

            if "namespaces" in sites[site]:
                for item in sites[site]["namespaces"].values():
                    if "loadbalancer" in item.keys():
                        for lb_type in item["loadbalancer"].keys():
                            lbs = lbs + len(item["loadbalancer"][lb_type].keys())

                for item in sites[site]["namespaces"].values():
                    if "origin_pools" in item.keys():
                        ops = ops + len(item["origin_pools"].keys())

                for item in sites[site]["namespaces"].values():
                    if "proxys" in item.keys():
                        proxies = proxies + len(item["proxys"].keys())

            table_data_services = [
                ('Count of LBs', lbs),
                ('Count of Origin pools', ops),
                ('Count of EFP', len(sites[site]["efp"].keys()) if "efp" in sites[site] else 0),
                ('Count of FPP', len(sites[site]["fpp"].keys()) if "fpp" in sites[site] else 0),
                ('Count of SMG', len(sites[site]["smg"].keys()) if "smg" in sites[site] else 0),
                ('Count of DCCG', len(sites[site]["dc_cluster_group"].keys()) if "dc_cluster_group" in sites[site] else 0),
                ('Count of Proxies', proxies),
                ('Count of Segments', len(sites[site]["segments"].keys()) if "segments" in sites[site] else 0),
                ('Count of BGP Policies', len(sites[site]["bgp"].keys()) if "bgp" in sites[site] else 0),
            ]

            ws_summary.append([f"Summary: {site}"] if title_prefix is None else [f"{title_prefix}: {site}"])
            ws_summary.merge_cells(f"A{ws_summary.max_row}:C{ws_summary.max_row}")
            for cell in ws_summary[ws_summary.max_row]:
                cell.fill = GREY_FILL
                cell.font = HEADER_FONT
                cell.alignment = LEFT_CENTER_ALIGNMENT

            ### Infrastructure Section
            ws_summary.append(["Infrastructure", ""])

            for cell in ws_summary[ws_summary.max_row]:
                cell.fill = GREY_FILL_SECTION
                cell.font = SECTION_FONT

            header = ["Item", "Value"]
            ws_summary.append(header)

            for cell in ws_summary[ws_summary.max_row]:
                cell.fill = GREY_FILL_SECTION
                if cell.column != 1:
                    cell.alignment = RIGHT_ALIGNMENT

            append_count = 0
            for item, value in table_data_infrastructure:
                ws_summary.append([item, value])
                append_count = append_count + 1

            max_row = ws_summary.max_row
            if ws_summary.max_row % 2 == 0:
                max_row = ws_summary.max_row + 1

            for idx in range(ws_summary.max_row - append_count, max_row):
                value_cell = ws_summary[idx][1]
                value_cell.alignment = RIGHT_ALIGNMENT

                # Check if the row number is EVEN
                if idx % 2 == 0:
                    # Apply the grey fill to every cell in the current row
                    for cell in ws_summary[idx]:
                        cell.fill = LIGHT_GREY_FILL

            ### Services Section
            ws_summary.append(["Services", ""])

            for cell in ws_summary[ws_summary.max_row]:
                cell.fill = GREY_FILL_SECTION
                cell.font = SECTION_FONT

            header = ["Item", "Value"]
            ws_summary.append(header)

            for cell in ws_summary[ws_summary.max_row]:
                cell.fill = GREY_FILL_SECTION
                if cell.column != 1:
                    cell.alignment = RIGHT_ALIGNMENT

            append_count = 0
            for item, value in table_data_services:
                ws_summary.append([item, value])
                append_count += 1

            for idx in range(ws_summary.max_row - append_count, ws_summary.max_row + 1):
                value_cell = ws_summary[idx][1]
                value_cell.alignment = RIGHT_ALIGNMENT

                # Check if the row number is EVEN
                if idx % 2 == 0:
                    # Apply the grey fill to every cell in the current row
                    for cell in ws_summary[idx]:
                        cell.fill = LIGHT_GREY_FILL

        with get_manager() as manager:
            with manager.counter(total=None, desc='Processing summary for', unit='sites') as pbar:
                for site, site_data in sites.items():
                    if self.must_break:
                        break
                    else:
                        if self.site:
                            if self.site == site:
                                self.must_break = True
                                process()
                                break
                        else:
                            process()

    def build_inventory_infrastructure(self, order: int = None, data: dict = None):
        """

        Parameters
        ----------
        data : dict
        order: int

        Returns
        -------

        """

        sites = data[c.SITES_KEY]
        site = None

        if self.site != "":
            site = sites[self.site]

        # WS Infrastructure Tab
        ws_infrastructure = self.wb.create_sheet("Infrastructure", order)
        ws_infrastructure.column_dimensions['A'].width = 30
        ws_infrastructure.column_dimensions['B'].width = 80
        ws_infrastructure.column_dimensions['C'].width = 80

        table_data_infrastructure = [
            ("Kind", site["kind"]),
            ("Provider Type", site["metadata"]["labels"]["ves.io/provider"] if "ves.io/provider" in site["metadata"]["labels"] else "Unknown"),
            ("Main Node Count", site["main_node_count"]),
            ("Worker Node Count", site["worker_node_count"] if "worker_node_count" in data else 0),
        ]

        if site["kind"] == c.F5XC_SITE_TYPE_SMS_V1 or site["kind"] == c.F5XC_SITE_TYPE_SMS_V2:
            table_data_infrastructure.append(("Labels", join_dict_items(site["sms"]["metadata"]["labels"])))
        else:
            table_data_infrastructure.append(("Labels", join_dict_items(site["legacy"]["metadata"]["labels"])))

        table_data_infrastructure.extend(
            [
                ("Node0 Hostname", site["nodes"]["node0"]["hostname"]),
                ("Node0 CPU Count", site["nodes"]["node0"]["hw_info"]["cpu"]["cpus"] if "hw_info" in site["nodes"]["node0"] else 0),
                ("Node0 CPU Model", site["nodes"]["node0"]["hw_info"]["cpu"]["model"] if "hw_info" in site["nodes"]["node0"] else 0),
                ("Node0 Memory Size (GB)", round(site["nodes"]["node0"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in site["nodes"]["node0"] else 0),
                ("Node0 Interface Count", len(site["nodes"]["node0"]["interfaces"]) if "interfaces" in site["nodes"]["node0"] else 0),
                ("Node0 OS Name", site["nodes"]["node0"]["hw_info"]["os"]["name"] if "hw_info" in site["nodes"]["node0"] else "None"),
                ("Node0 OS Version", site["nodes"]["node0"]["hw_info"]["os"]["version"] if "hw_info" in site["nodes"]["node0"] else "None")
            ]
        )

        if "hw_info" in site["nodes"]["node0"] and "hw_info" in site["nodes"]["node0"]:
            for storage_source in site["nodes"]["node0"]["hw_info"]["storage"]:
                source_node0_storage_size = storage_source["size_gb"]
                table_data_infrastructure.append((f"Node0 Storage {storage_source["name"]} Size (GB)", source_node0_storage_size))

        if site["main_node_count"] > 1 and site["main_node_count"] > 1:
            table_data_infrastructure.extend(
                [
                    ("Node1 Hostname", site["nodes"]["node1"]["hostname"]),
                    ("Node1 CPU Count", site["nodes"]["node1"]["hw_info"]["cpu"]["cpus"] if "hw_info" in site["nodes"]["node1"] else 0),
                    ("Node1 CPU Model", site["nodes"]["node1"]["hw_info"]["cpu"]["model"] if "hw_info" in site["nodes"]["node1"] else 0),
                    ("Node1 Memory Size (GB)", round(site["nodes"]["node1"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in site["nodes"]["node1"] else 0),
                    ("Node1 Interface Count", len(site["nodes"]["node1"]["interfaces"]) if "interfaces" in site["nodes"]["node1"] else 0),
                    ("Node1 OS Name", site["nodes"]["node1"]["hw_info"]["os"]["name"] if "hw_info" in site["nodes"]["node1"] else "None"),
                    ("Node1 OS Version", site["nodes"]["node1"]["hw_info"]["os"]["version"] if "hw_info" in site["nodes"]["node1"] else "None"),
                ]
            )

            if "hw_info" in site["nodes"]["node1"]:
                for storage_site in site["nodes"]["node1"]["hw_info"]["storage"]:
                    site_node1_storage_size = storage_site["size_gb"]
                    table_data_infrastructure.append((f"Node1 Storage {storage_site["name"]} Size (GB)", site_node1_storage_size))

            table_data_infrastructure.extend(
                [
                    ("Node2 Hostname", site["nodes"]["node2"]["hostname"]),
                    ("Node2 CPU Count", site["nodes"]["node2"]["hw_info"]["cpu"]["cpus"] if "hw_info" in site["nodes"]["node2"] else 0),
                    ("Node2 CPU Model", site["nodes"]["node2"]["hw_info"]["cpu"]["model"] if "hw_info" in site["nodes"]["node2"] else 0),
                    ("Node2 Memory Size (GB)", round(site["nodes"]["node2"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in site["nodes"]["node2"] else 0),
                    ("Node2 Interface Count", len(site["nodes"]["node2"]["interfaces"]) if "interfaces" in site["nodes"]["node2"] else 0),
                    ("Node2 OS Name", site["nodes"]["node2"]["hw_info"]["os"]["name"] if "hw_info" in site["nodes"]["node2"] else "None"),
                    ("Node2 OS Version", site["nodes"]["node2"]["hw_info"]["os"]["version"] if "hw_info" in site["nodes"]["node2"] else "None")
                ]
            )

            if "hw_info" in site["nodes"]["node2"]:
                for storage_site in site["nodes"]["node2"]["hw_info"]["storage"]:
                    site_node1_storage_size = storage_site["size_gb"]
                    table_data_infrastructure.append((f"Node2 Storage {storage_site["name"]} Size (GB)", site_node1_storage_size))

        # Node0 Interface computation
        site_node0_interfaces = list()
        if site["kind"] == c.F5XC_SITE_TYPE_SMS_V1:
            for interface in site["nodes"]["node0"]["interfaces"]:
                interface_details = dict()
                if "dedicated_interface" in interface.keys():
                    interface_details["is_primary"] = "true" if "is_primary" in interface["dedicated_interface"].keys() else "false"
                    interface_details["device_name"] = interface["dedicated_interface"]["device"]
                    interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                    interface_details["interface_type"] = "dedicated_interface"
                    _interfaces = [f"Node0", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                    site_node0_interfaces.append(_interfaces)
                if "ethernet_interface" in interface.keys():
                    interface_details["mtu"] = interface["ethernet_interface"]["mtu"]
                    interface_details["is_primary"] = True if "is_primary" in interface["ethernet_interface"].keys() else False
                    interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                    interface_details["device_name"] = interface["ethernet_interface"]["device"]
                    interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                    if "dhcp_server" in interface["ethernet_interface"].keys():
                        network_prefixes = list()
                        for network in interface["ethernet_interface"]["dhcp_server"]["dhcp_networks"]:
                            network_prefixes.append(network["network_prefix"])
                        interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                    interface_details["interface_type"] = "ethernet_interface"
                    interface_details["segment_network"] = interface["ethernet_interface"]["segment_network"]["name"] if "segment_network" in interface[
                        "ethernet_interface"].keys() else "None"
                    _interface = ["Node0", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                    site_node0_interfaces.append(_interface)
        elif  site["kind"] == c.F5XC_SITE_TYPE_SMS_V2:
            for interface in site["nodes"]["node0"]["interfaces"]:
                interface_details = dict()
                interface_details["mtu"] = interface["mtu"] if "mtu" in interface else "None"
                interface_details["is_primary"] = interface["is_primary"] if "is_primary" in interface else "None"
                interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                interface_details["is_management"] =  interface["is_management"] if "is_management" in interface else "None"
                interface_details["site_local_network"] = "True" if "site_local_network" in interface["network_option"] else "None"
                interface_details["site_local_inside_network"] = "True" if "site_local_inside_network" in interface["network_option"] else "None"
                interface_details["dhcp_client"] = "True" if "dhcp_client" in interface else "False"
                interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                interface_details["segment_network"] = interface["network_option"]["segment_network"]["name"] if "segment_network" in interface["network_option"].keys() else "None"
                _interface = ["Node0", f"{interface["name"]}", join_dict_items(interface_details)]
                if "dhcp_server" in interface.keys():
                    network_prefixes = list()
                    for network in interface["dhcp_server"]["dhcp_networks"]:
                        network_prefixes.append(network["network_prefix"])
                    interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                if "ethernet_interface" in interface.keys():
                    interface_details["device_name"] = interface["ethernet_interface"]["device"]
                    interface_details["interface_type"] = "ethernet_interface"
                    interface_details["mac"] = interface["ethernet_interface"]["mac"] if "ethernet_interface" in interface["ethernet_interface"].keys() else "None"
                site_node0_interfaces.append(_interface)
        else:
            # Legacy sites
            for interface_name, interface_attrs in site["nodes"]["node0"]["interfaces"].items():
                interface_details = dict()
                interface_details["device_name"] = interface_name
                interface_details["ipv4"] = interface_attrs["subnet_param"]["ipv4"] if "subnet_param" in interface_attrs else None
                interface_details["ipv6"] = interface_attrs["subnet_param"]["ipv6"] if "subnet_param" in interface_attrs else None
                interface_details["existing_subnet_id"] = interface_attrs["existing_subnet_id"] if "existing_subnet_id" in interface_attrs else None
                _interface = [f"Node0", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                site_node0_interfaces.append(_interface)

        table_data_infrastructure_interfaces = []
        for site_interface in site_node0_interfaces:
            table_data_infrastructure_interfaces.append(tuple(site_interface))

        if site["main_node_count"] > 1:
            # Node1 Interface computation
            site_node1_interfaces = list()
            if site["kind"] == c.F5XC_SITE_TYPE_SMS_V1:
                for interface in site["nodes"]["node1"]["interfaces"]:
                    interface_details = dict()
                    if "dedicated_interface" in interface.keys():
                        interface_details["is_primary"] = "true" if "is_primary" in interface["dedicated_interface"].keys() else "false"
                        interface_details["device_name"] = interface["dedicated_interface"]["device"]
                        interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                        interface_details["interface_type"] = "dedicated_interface"
                        _interfaces = [f"Node1", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                        site_node1_interfaces.append(_interfaces)
                    if "ethernet_interface" in interface.keys():
                        interface_details["mtu"] = interface["ethernet_interface"]["mtu"]
                        interface_details["is_primary"] = True if "is_primary" in interface["ethernet_interface"].keys() else False
                        interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                        interface_details["device_name"] = interface["ethernet_interface"]["device"]
                        interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                        if "dhcp_server" in interface["ethernet_interface"].keys():
                            network_prefixes = list()
                            for network in interface["ethernet_interface"]["dhcp_server"]["dhcp_networks"]:
                                network_prefixes.append(network["network_prefix"])
                            interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                        interface_details["interface_type"] = "ethernet_interface"
                        interface_details["segment_network"] = interface["ethernet_interface"]["segment_network"]["name"] if "segment_network" in interface[
                            "ethernet_interface"].keys() else "None"
                        _interface = ["Node1", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                        site_node1_interfaces.append(_interface)
            elif site["kind"] == c.F5XC_SITE_TYPE_SMS_V2:
                for interface in site["nodes"]["node1"]["interfaces"]:
                    interface_details = dict()
                    interface_details["mtu"] = interface["mtu"] if "mtu" in interface else "None"
                    interface_details["is_primary"] = interface["is_primary"] if "is_primary" in interface else "None"
                    interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                    interface_details["is_management"] = interface["is_management"] if "is_management" in interface else "None"
                    interface_details["site_local_network"] = "True" if "site_local_network" in interface["network_option"] else "None"
                    interface_details["site_local_inside_network"] = "True" if "site_local_inside_network" in interface["network_option"] else "None"
                    interface_details["dhcp_client"] = "True" if "dhcp_client" in interface else "False"
                    interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                    interface_details["segment_network"] = interface["network_option"]["segment_network"]["name"] if "segment_network" in interface[
                        "network_option"].keys() else "None"
                    if "dhcp_server" in interface.keys():
                        network_prefixes = list()
                        for network in interface["dhcp_server"]["dhcp_networks"]:
                            network_prefixes.append(network["network_prefix"])
                        interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                    if "ethernet_interface" in interface.keys():
                        interface_details["device_name"] = interface["ethernet_interface"]["device"]
                        interface_details["interface_type"] = "ethernet_interface"
                        interface_details["mac"] = interface["ethernet_interface"]["mac"] if "ethernet_interface" in interface["ethernet_interface"].keys() else "None"
                    _interface = ["Node0", f"{interface["name"]}", join_dict_items(interface_details)]
                    site_node1_interfaces.append(_interface)
            else:
                # Legacy sites
                for interface_name, interface_attrs in site["nodes"]["node1"]["interfaces"].items():
                    interface_details = dict()
                    interface_details["device_name"] = interface_name
                    interface_details["ipv4"] = interface_attrs["subnet_param"]["ipv4"] if "subnet_param" in interface_attrs else None
                    interface_details["ipv6"] = interface_attrs["subnet_param"]["ipv6"] if "subnet_param" in interface_attrs else None
                    interface_details["existing_subnet_id"] = interface_attrs["existing_subnet_id"] if "existing_subnet_id" in interface_attrs else None
                    _interface = [f"Node1", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                    site_node1_interfaces.append(_interface)

            for site_interface in site_node1_interfaces:
                table_data_infrastructure_interfaces.append(tuple(site_interface))

            # Node2 Interface computation
            site_node2_interfaces = list()
            if site["kind"] == c.F5XC_SITE_TYPE_SMS_V1:
                for interface in site["nodes"]["node2"]["interfaces"]:
                    interface_details = dict()
                    if "dedicated_interface" in interface.keys():
                        interface_details["is_primary"] = "true" if "is_primary" in interface["dedicated_interface"].keys() else "false"
                        interface_details["device_name"] = interface["dedicated_interface"]["device"]
                        interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                        interface_details["interface_type"] = "dedicated_interface"
                        _interfaces = [f"Node2", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                        site_node2_interfaces.append(_interfaces)
                    if "ethernet_interface" in interface.keys():
                        interface_details["mtu"] = interface["ethernet_interface"]["mtu"]
                        interface_details["is_primary"] = True if "is_primary" in interface["ethernet_interface"].keys() else False
                        interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                        interface_details["device_name"] = interface["ethernet_interface"]["device"]
                        interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                        if "dhcp_server" in interface["ethernet_interface"].keys():
                            network_prefixes = list()
                            for network in interface["ethernet_interface"]["dhcp_server"]["dhcp_networks"]:
                                network_prefixes.append(network["network_prefix"])
                            interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                        interface_details["interface_type"] = "ethernet_interface"
                        interface_details["segment_network"] = interface["ethernet_interface"]["segment_network"]["name"] if "segment_network" in interface[
                            "ethernet_interface"].keys() else "None"
                        _interface = ["Node2", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                        site_node2_interfaces.append(_interface)
            elif site["kind"] == c.F5XC_SITE_TYPE_SMS_V2:
                for interface in site["nodes"]["node2"]["interfaces"]:
                    interface_details = dict()
                    interface_details["mtu"] = interface["mtu"] if "mtu" in interface else "None"
                    interface_details["is_primary"] = interface["is_primary"] if "is_primary" in interface else "None"
                    interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                    interface_details["is_management"] = interface["is_management"] if "is_management" in interface else "None"
                    interface_details["site_local_network"] = "True" if "site_local_network" in interface["network_option"] else "None"
                    interface_details["site_local_inside_network"] = "True" if "site_local_inside_network" in interface["network_option"] else "None"
                    interface_details["dhcp_client"] = "True" if "dhcp_client" in interface else "False"
                    interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                    interface_details["segment_network"] = interface["network_option"]["segment_network"]["name"] if "segment_network" in interface[
                        "network_option"].keys() else "None"
                    if "dhcp_server" in interface.keys():
                        network_prefixes = list()
                        for network in interface["dhcp_server"]["dhcp_networks"]:
                            network_prefixes.append(network["network_prefix"])
                        interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                    if "ethernet_interface" in interface.keys():
                        interface_details["device_name"] = interface["ethernet_interface"]["device"]
                        interface_details["interface_type"] = "ethernet_interface"
                        interface_details["mac"] = interface["ethernet_interface"]["mac"] if "ethernet_interface" in interface["ethernet_interface"].keys() else "None"
                    _interface = ["Node2", f"{interface["name"]}", join_dict_items(interface_details)]
                    site_node2_interfaces.append(_interface)
            else:
                # Legacy sites
                for interface_name, interface_attrs in site["nodes"]["node2"]["interfaces"].items():
                    interface_details = dict()
                    interface_details["device_name"] = interface_name
                    interface_details["ipv4"] = interface_attrs["subnet_param"]["ipv4"] if "subnet_param" in interface_attrs else None
                    interface_details["ipv6"] = interface_attrs["subnet_param"]["ipv6"] if "subnet_param" in interface_attrs else None
                    interface_details["existing_subnet_id"] = interface_attrs["existing_subnet_id"] if "existing_subnet_id" in interface_attrs else None
                    _interface = [f"Node2", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                    site_node2_interfaces.append(_interface)

            for site_interface in site_node2_interfaces:
                table_data_infrastructure_interfaces.append(tuple(site_interface))

        ws_infrastructure.append([f"Infrastructure: {site["metadata"]["name"]}"])
        ws_infrastructure.merge_cells(f"A{ws_infrastructure.max_row}:C{ws_infrastructure.max_row}")
        for cell in ws_infrastructure[ws_infrastructure.max_row]:
            cell.fill = GREY_FILL
            cell.font = HEADER_FONT
            cell.alignment = LEFT_CENTER_ALIGNMENT

        # Nodes section
        ws_infrastructure.append(["Hardware/Software", ""])

        for cell in ws_infrastructure[ws_infrastructure.max_row]:
            cell.fill = GREY_FILL_SECTION
            cell.font = SECTION_FONT

        header = ["Item", "Value"]
        ws_infrastructure.append(header)

        for cell in ws_infrastructure[ws_infrastructure.max_row]:
            cell.fill = GREY_FILL_SECTION
            if cell.column != 1:
                cell.alignment = LEFT_ALIGNMENT

        append_count = 0
        for item, value in table_data_infrastructure:
            ws_infrastructure.append([item, value])
            append_count = append_count + 1

        start = 4
        end = ws_infrastructure.max_row + 1

        for row_num in range(start, end):
            value_cell_site = ws_infrastructure[row_num][1]
            value_cell_site.alignment = RIGHT_ALIGNMENT
            value_cell_site.font = Font(color='FF8B0000', bold=False)

            # Check if the row number is EVEN
            if row_num % 2 == 0:
                # Apply the grey fill to every cell in the current row
                for cell in ws_infrastructure[row_num]:
                    cell.fill = LIGHT_GREY_FILL

        # Nodes interfaces section
        ws_infrastructure.append(["Interfaces", ""])

        for cell in ws_infrastructure[ws_infrastructure.max_row]:
            cell.fill = GREY_FILL_SECTION
            cell.font = SECTION_FONT

        header = ["Node", "Interface", "Values"]
        ws_infrastructure.append(header)

        for cell in ws_infrastructure[ws_infrastructure.max_row]:
            cell.fill = GREY_FILL_SECTION
            if cell.column != 1:
                cell.alignment = LEFT_ALIGNMENT

        append_count_interface = 0
        for node, source_iface, source_iface_values in table_data_infrastructure_interfaces:
            ws_infrastructure.append([node, source_iface, source_iface_values])
            append_count_interface = append_count_interface + 1

        start = ws_infrastructure.max_row - append_count_interface + 1
        end = ws_infrastructure.max_row + 1

        for row_num in range(start, end):
            node_cell = ws_infrastructure[row_num][0]
            interface_cell_site = ws_infrastructure[row_num][1]
            value_cell_site = ws_infrastructure[row_num][2]

            node_cell.alignment = Alignment(horizontal="center", vertical='center')
            interface_cell_site.alignment = Alignment(horizontal="center", vertical='center')
            value_cell_site.font = Font(color='FF8B0000', bold=False)

            # Check if the row number is EVEN
            if row_num % 2 == 0:
                # Apply the grey fill to every cell in the current row
                for cell in ws_infrastructure[row_num]:
                    cell.fill = LIGHT_GREY_FILL

            max_height_for_row = 0

            for col_letter in ["C"]:
                cell = ws_infrastructure[f'{col_letter}{row_num}']

                if cell.value == "N/A":
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')

                cell_c = ws_infrastructure[f'C{row_num}']
                required_height = len(str(cell_c.value)) + 10

                if required_height is not None and required_height > max_height_for_row:
                    max_height_for_row = required_height

                if max_height_for_row > 0:
                    ws_infrastructure.row_dimensions[row_num].height = max_height_for_row

    def build_inventory_service(self, order: int = None, data: dict = None):
        """

        Parameters
        ----------
        data: dict
        order: int

        Returns
        -------

        """

        sites = data[c.SITES_KEY]
        site = None

        if self.site != "":
            site = sites[self.site]

        # WS Services Tab
        ws_services = self.wb.create_sheet("Services", order)
        ws_services.column_dimensions['A'].width = 25
        ws_services.column_dimensions['B'].width = 60
        ws_services.column_dimensions['C'].width = 20

        site_ns = list()
        site_lbs = list()
        site_ops = list()
        site_proxies = list()

        if "namespaces" in site:
            for namespace in site["namespaces"]:
                site_ns.append(namespace)
            for site_item in site["namespaces"].values():
                if "loadbalancer" in site_item.keys():
                    for source_lb_type in site_item["loadbalancer"].keys():
                        site_lbs.append(list(site_item["loadbalancer"][source_lb_type].keys())[0])

                if "proxys" in site_item.keys():
                    for source_proxy_type in site_item["proxys"].keys():
                        site_proxies.append(site_item["proxys"][source_proxy_type]["metadata"]["name"])

            for site_item in site["namespaces"].values():
                if "origin_pools" in site_item.keys():
                    site_ops.append(list(site_item["origin_pools"].keys())[0])

        table_data_services = [
            ('NS', "\n".join(site_ns) if len(site_ns) > 0 else "None"),
            ('LB', "\n".join(site_lbs) if len(site_lbs) > 0 else "None"),
            ('OP', "\n".join(site_ops) if len(site_ops) > 0 else "None"),
            ('EFP', "\n".join(site["efp"].keys()) if "efp" in site else "None"),
            ('FPP', "\n".join(site["fpp"].keys()) if "fpp" in site else "None"),
            ('SMG', "\n".join(site["smg"].keys()) if len(site["smg"]) > 0 else "None"),
            ('DCCG', "\n".join(site["dc_cluster_group"].keys()) if "dc_cluster_group" in site else "None"),
            ('Proxies', "\n".join(site_proxies) if len(site_proxies) > 0 else "None"),
            ('Segments', "\n".join(site["segments"].keys()) if "segments" in site else "None"),
            ('BGP Policies', "\n".join(site["bgp"].keys()) if "bgp" in site else "None"),
            ('Virtual Sites', "\n".join(site["vsites"]) if len(site["vsites"]) > 0 else "None"),
        ]

        ws_services.append([f"Services: {site["metadata"]["name"]}"])
        ws_services.merge_cells(f"A{ws_services.max_row}:C{ws_services.max_row}")
        for cell in ws_services[ws_services.max_row]:
            cell.fill = GREY_FILL
            cell.font = HEADER_FONT
            cell.alignment = LEFT_CENTER_ALIGNMENT

        ws_services.append(["Services", ""])

        for cell in ws_services[ws_services.max_row]:
            cell.fill = GREY_FILL_SECTION
            cell.font = SECTION_FONT

        header = ["Item", "Value",]
        ws_services.append(header)

        for cell in ws_services[ws_services.max_row]:
            cell.fill = GREY_FILL_SECTION
            if cell.column != 1:
                cell.alignment = LEFT_ALIGNMENT

        append_count = 0
        for item, source_value in table_data_services:
            ws_services.append([item, source_value])
            append_count += 1

        for row_num in range(ws_services.max_row - append_count + 1, ws_services.max_row + 1):
            value_cell = ws_services[row_num][1]
            value_cell.alignment = LEFT_ALIGNMENT
            value_cell.font = Font(color='FF8B0000', bold=False)

            # Check if the row number is EVEN
            if row_num % 2 == 0:
                # Apply the grey fill to every cell in the current row
                for cell in ws_services[row_num]:
                    cell.fill = LIGHT_GREY_FILL

            max_height_for_row = 0

            for col_letter in ["B", "C"]:
                cell = ws_services[f'{col_letter}{row_num}']
                cell.alignment = Alignment(wrap_text=True, vertical='top')

                cell_b = ws_services[f'B{row_num}']
                required_height = len(str(cell_b.value)) + 10

                if required_height is not None and required_height > max_height_for_row:
                    max_height_for_row = required_height

                if max_height_for_row > 0:
                    ws_services.row_dimensions[row_num].height = max_height_for_row

        for col_letter in ["A"]:
            for row_num in range(4, ws_services.max_row + 1):
                cell = ws_services[f'{col_letter}{row_num}']
                cell.alignment = Alignment(horizontal='center', vertical='center')

    def build_inventory(self, data: dict = None) -> bool:
        """

        Parameters
        ----------
        data: dict

        Returns
        -------

        """

        if self.site == "":
            self.logger.info(f"Building inventory failed. Site name required. use -s option to provide site name")

            return False
        else:
            self.build_inventory_summary(0, data)
            self.build_inventory_infrastructure(1, data)
            self.build_inventory_service(2, data)

            return True

    def build_compare_summary(self, order: int = None, data_source: dict = None, data_target: dict = None):
        """

        Parameters
        ----------
        order: int
        data_source: dict
        data_target: dict

        Returns
        -------

        """

        # WS Summary Comparison Tab
        ws_summary = self.wb.create_sheet("Summary", order)
        ws_summary.column_dimensions['A'].width = 25
        ws_summary.column_dimensions['B'].width = 30
        ws_summary.column_dimensions['C'].width = 30

        table_data_infrastructure = [
            ("Kind", data_source["kind"], data_target["kind"]),
            ("Provider Type", data_source["metadata"]["labels"]["ves.io/provider"] if "ves.io/provider" in data_source["metadata"]["labels"] else "Unknown",
             data_target["metadata"]["labels"]["ves.io/provider"] if "ves.io/provider" in data_target["metadata"]["labels"] else "Unknown"),
            ("Main Node Count", data_source["main_node_count"], data_target["main_node_count"]),
            ("Worker Node Count", data_source["worker_node_count"] if "worker_node_count" in data_source else 0,
             data_target["worker_node_count"] if "worker_node_count" in data_target else 0),
            ("Node0 CPU Count", data_source["nodes"]["node0"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_source["nodes"]["node0"] else 0,
             data_target["nodes"]["node0"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_target["nodes"]["node0"] else 0),
            ("Node0 Memory Size (GB)", round(data_source["nodes"]["node0"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_source["nodes"]["node0"] else 0,
             round(data_target["nodes"]["node0"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_target["nodes"]["node0"] else 0),
            ("Node0 Interface Count", len(data_source["nodes"]["node0"]["interfaces"]) if "interfaces" in data_source["nodes"]["node0"] else 0,
             len(data_target["nodes"]["node0"]["interfaces"]) if "interfaces" in data_target["nodes"]["node0"] else 0)
        ]

        if "hw_info" in data_source["nodes"]["node0"] and "hw_info" in data_target["nodes"]["node0"]:
            for storage_source, storage_target in zip(data_source["nodes"]["node0"]["hw_info"]["storage"], data_target["nodes"]["node0"]["hw_info"]["storage"]):
                source_node0_storage_size = storage_source["size_gb"]
                target_node0_storage_size = storage_target["size_gb"]
                table_data_infrastructure.append((f"Node0 Storage {storage_source["name"]} Size (GB)", source_node0_storage_size, target_node0_storage_size))

        if data_source["main_node_count"] > 1 and data_target["main_node_count"] > 1:
            table_data_infrastructure.extend(
                [
                    ("Node1 CPU Count", data_source["nodes"]["node1"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_source["nodes"]["node1"] else 0,
                     data_target["nodes"]["node1"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_target["nodes"]["node1"] else 0),
                    ("Node1 Memory Size (GB)", round(data_source["nodes"]["node1"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_source["nodes"]["node1"] else 0,
                     round(data_target["nodes"]["node1"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_target["nodes"]["node1"] else 0),
                    ("Node1 Interface Count", len(data_source["nodes"]["node1"]["interfaces"]) if "interfaces" in data_source["nodes"]["node1"] else 0,
                     len(data_target["nodes"]["node1"]["interfaces"]) if "interfaces" in data_target["nodes"]["node1"] else 0),
                ]
            )

            if "hw_info" in data_source["nodes"]["node1"]:
                for storage_source, storage_target in zip(data_source["nodes"]["node1"]["hw_info"]["storage"], data_target["nodes"]["node1"]["hw_info"]["storage"]):
                    source_node1_storage_size = storage_source["size_gb"]
                    target_node1_storage_size = storage_target["size_gb"]
                    table_data_infrastructure.append((f"Node1 Storage {storage_source["name"]} Size (GB)", source_node1_storage_size, target_node1_storage_size))

            table_data_infrastructure.extend(
                [
                    ("Node2 CPU Count", data_source["nodes"]["node2"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_source["nodes"]["node2"] else 0,
                     data_target["nodes"]["node2"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_target["nodes"]["node2"] else 0),
                    ("Node2 Memory Size (GB)", round(data_source["nodes"]["node2"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_source["nodes"]["node2"] else 0,
                     round(data_target["nodes"]["node2"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_target["nodes"]["node2"] else 0),
                    ("Node2 Interface Count", len(data_source["nodes"]["node2"]["interfaces"]) if "interfaces" in data_source["nodes"]["node2"] else 0,
                     len(data_target["nodes"]["node2"]["interfaces"]) if "interfaces" in data_target["nodes"]["node2"] else 0),
                ]
            )

            if "hw_info" in data_source["nodes"]["node2"]:
                for storage_source, storage_target in zip(data_source["nodes"]["node2"]["hw_info"]["storage"], data_target["nodes"]["node2"]["hw_info"]["storage"]):
                    source_node2_storage_size = storage_source["size_gb"]
                    target_node2_storage_size = storage_target["size_gb"]
                    table_data_infrastructure.append((f"Node2 Storage {storage_source["name"]} Size (GB)", source_node2_storage_size, target_node2_storage_size))

        source_lbs = 0
        source_ops = 0
        target_lbs = 0
        target_ops = 0
        source_proxies = 0
        target_proxies = 0

        if "namespaces" in data_source:
            for source_item in data_source["namespaces"].values():
                if "loadbalancer" in source_item.keys():
                    for source_lb_type in source_item["loadbalancer"].keys():
                        source_lbs = source_lbs + len(source_item["loadbalancer"][source_lb_type].keys())
                if "proxys" in source_item.keys():
                    for source_proxy_type in source_item["proxys"].keys():
                        source_proxies = source_proxies + len(source_item["proxys"][source_proxy_type].keys())

            for source_item in data_source["namespaces"].values():
                if "origin_pools" in source_item.keys():
                    source_ops = source_ops + len(source_item["origin_pools"].keys())

        if "namespaces" in data_target:
            for target_item in data_target["namespaces"].values():
                if "loadbalancer" in target_item.keys():
                    for target_lb_type in target_item["loadbalancer"].keys():
                        target_lbs = target_lbs + len(target_item["loadbalancer"][target_lb_type].keys())
                if "proxys" in target_item.keys():
                    for target_proxy_type in target_item["proxys"].keys():
                        target_proxies = target_proxies + len(target_item["proxys"][target_proxy_type].keys())

            for target_item in data_target["namespaces"].values():
                if "origin_pools" in target_item.keys():
                    target_ops = target_ops + len(target_item["origin_pools"].keys())

        table_data_services = [
            ('Count of LBs', source_lbs, target_lbs),
            ('Count of Origin pools', source_ops, target_ops),
            ('Count of EFP', len(data_source["efp"].keys()) if "efp" in data_source else 0, len(data_target["efp"].keys()) if "efp" in data_target else 0),
            ('Count of FPP', len(data_source["fpp"].keys()) if "fpp" in data_source else 0, len(data_target["fpp"].keys()) if "fpp" in data_target else 0),
            ('Count of SMG', len(data_source["smg"].keys()) if "smg" in data_source else 0, len(data_target["smg"].keys()) if "smg" in data_target else 0),
            ('Count of DCCG', len(data_source["dc_cluster_group"].keys()) if "dc_cluster_group" in data_source else 0,
             len(data_target["dc_cluster_group"].keys()) if "dc_cluster_group" in data_target else 0),
            ('Count of Proxies', source_proxies, target_proxies),
            ('Count of Segments', len(data_source["segments"].keys()) if "segments" in data_source else 0, len(data_target["segments"].keys()) if "segments" in data_target else 0),
            ('Count of BGP Policies', len(data_source["bgp"].keys()) if "bgp" in data_source else 0, len(data_target["bgp"].keys()) if "bgp" in data_target else 0),
            ('Count of Virtual Sites', len(data_source["vsites"]), len(data_target["vsites"]))
        ]

        ws_summary.append([f"Summary comparison: {data_source["metadata"]["name"]} with {data_target["metadata"]["name"]}"])
        ws_summary.merge_cells(f"A{ws_summary.max_row}:F{ws_summary.max_row}")
        for cell in ws_summary[ws_summary.max_row]:
            cell.fill = GREY_FILL
            cell.font = HEADER_FONT
            cell.alignment = LEFT_CENTER_ALIGNMENT

        ### Infrastructure Section
        ws_summary.append(["Infrastructure", ""])

        for cell in ws_summary[ws_summary.max_row]:
            cell.fill = GREY_FILL_SECTION
            cell.font = SECTION_FONT

        header = ["Item", "Source", "Target"]
        ws_summary.append(header)

        for cell in ws_summary[ws_summary.max_row]:
            cell.fill = GREY_FILL_SECTION
            if cell.column != 1:
                cell.alignment = RIGHT_ALIGNMENT

        append_count = 0
        for item, source, target in table_data_infrastructure:
            ws_summary.append([item, source, target])
            append_count = append_count + 1

        start = 4
        end = ws_summary.max_row + 1

        for idx in range(start, end):
            value_cell_source = ws_summary[idx][1]
            value_cell_target = ws_summary[idx][2]
            value_cell_source.alignment = RIGHT_ALIGNMENT
            value_cell_target.alignment = RIGHT_ALIGNMENT

            if value_cell_source.value != value_cell_target.value:
                value_cell_source.font = Font(color='FFFF0000', bold=False)
                value_cell_target.font = Font(color='FFFF0000', bold=False)

            # Check if the row number is EVEN
            if idx % 2 == 0:
                # Apply the grey fill to every cell in the current row
                for cell in ws_summary[idx]:
                    cell.fill = LIGHT_GREY_FILL

        ### Services Section
        ws_summary.append(["Services", ""])

        for cell in ws_summary[ws_summary.max_row]:
            cell.fill = GREY_FILL_SECTION
            cell.font = SECTION_FONT

        header = ["Item", "Source", "Target"]
        ws_summary.append(header)

        for cell in ws_summary[ws_summary.max_row]:
            cell.fill = GREY_FILL_SECTION
            if cell.column != 1:
                cell.alignment = RIGHT_ALIGNMENT

        append_count = 0
        for item, source_value, target_value in table_data_services:
            ws_summary.append([item, source_value, target_value])
            append_count += 1

        for idx in range(ws_summary.max_row - append_count + 1, ws_summary.max_row + 1):
            value_cell_source = ws_summary[idx][1]
            value_cell_target = ws_summary[idx][2]
            value_cell_source.alignment = RIGHT_ALIGNMENT
            value_cell_target.alignment = RIGHT_ALIGNMENT

            if value_cell_source.value != value_cell_target.value:
                value_cell_source.font = Font(color='FFFF0000', bold=False)
                value_cell_target.font = Font(color='FFFF0000', bold=False)

            # Check if the row number is EVEN
            if idx % 2 == 0:
                # Apply the grey fill to every cell in the current row
                for cell in ws_summary[idx]:
                    cell.fill = LIGHT_GREY_FILL

    def build_compare_infrastructure(self, order: int = None, data_source: dict = None, data_target: dict = None):
        """

        Parameters
        ----------
        data_source
        data_target
        order: int

        Returns
        -------

        """

        # WS Infrastructure Tab
        ws_infrastructure = self.wb.create_sheet("Infrastructure", order)
        ws_infrastructure.column_dimensions['A'].width = 30
        ws_infrastructure.column_dimensions['B'].width = 80
        ws_infrastructure.column_dimensions['C'].width = 80
        ws_infrastructure.column_dimensions['D'].width = 20
        ws_infrastructure.column_dimensions['E'].width = 60

        table_data_infrastructure = [
            ("Kind", data_source["kind"], data_target["kind"]),
            ("Provider Type", data_source["metadata"]["labels"]["ves.io/provider"] if "ves.io/provider" in data_source["metadata"]["labels"] else "Unknown",
             data_target["metadata"]["labels"]["ves.io/provider"] if "ves.io/provider" in data_target["metadata"]["labels"] else "Unknown"),
            ("Main Node Count", data_source["main_node_count"], data_target["main_node_count"]),
            ("Worker Node Count", data_source["worker_node_count"] if "worker_node_count" in data_source else 0,
             data_target["worker_node_count"] if "worker_node_count" in data_target else 0),
        ]

        if data_source["kind"] == c.F5XC_SITE_TYPE_SMS_V1:
            table_data_infrastructure.append(("Labels", join_dict_items(data_source["sms"]["metadata"]["labels"]), join_dict_items(data_target["sms"]["metadata"]["labels"])))
        else:
            table_data_infrastructure.append(("Labels", join_dict_items(data_source["legacy"]["metadata"]["labels"]), join_dict_items(data_target["sms"]["metadata"]["labels"])))

        table_data_infrastructure.extend(
            [
                ("Node0 Hostname", data_source["nodes"]["node0"]["hostname"], data_target["nodes"]["node0"]["hostname"]),
                ("Node0 CPU Count", data_source["nodes"]["node0"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_source["nodes"]["node0"] else 0,
                 data_target["nodes"]["node0"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_target["nodes"]["node0"] else "None"),
                ("Node0 CPU Model", data_source["nodes"]["node0"]["hw_info"]["cpu"]["model"] if "hw_info" in data_source["nodes"]["node0"] else 0,
                 data_target["nodes"]["node0"]["hw_info"]["cpu"]["model"] if "hw_info" in data_target["nodes"]["node0"] else "None"),
                ("Node0 Memory Size (GB)", round(data_source["nodes"]["node0"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_source["nodes"]["node0"] else 0,
                 round(data_target["nodes"]["node0"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_target["nodes"]["node0"] else 0),
                ("Node0 Interface Count", len(data_source["nodes"]["node0"]["interfaces"]) if "interfaces" in data_source["nodes"]["node0"] else 0,
                 len(data_target["nodes"]["node0"]["interfaces"]) if "interfaces" in data_target["nodes"]["node0"] else 0),
                ("Node0 OS Name", data_source["nodes"]["node0"]["hw_info"]["os"]["name"] if "hw_info" in data_source["nodes"]["node0"] else "None",
                 data_target["nodes"]["node0"]["hw_info"]["os"]["name"] if "hw_info" in data_target["nodes"]["node0"] else "None"),
                ("Node0 OS Version", data_source["nodes"]["node0"]["hw_info"]["os"]["version"] if "hw_info" in data_source["nodes"]["node0"] else "None",
                 data_target["nodes"]["node0"]["hw_info"]["os"]["version"] if "hw_info" in data_target["nodes"]["node0"] else "None"),
            ]
        )

        if "hw_info" in data_source["nodes"]["node0"] and "hw_info" in data_target["nodes"]["node0"]:
            for storage_source, storage_target in zip(data_source["nodes"]["node0"]["hw_info"]["storage"], data_target["nodes"]["node0"]["hw_info"]["storage"]):
                source_node0_storage_size = storage_source["size_gb"]
                target_node0_storage_size = storage_target["size_gb"]
                table_data_infrastructure.append((f"Node0 Storage {storage_source["name"]} Size (GB)", source_node0_storage_size, target_node0_storage_size))

        if data_source["main_node_count"] > 1 and data_target["main_node_count"] > 1:
            table_data_infrastructure.extend(
                [
                    ("Node1 Hostname", data_source["nodes"]["node1"]["hostname"], data_target["nodes"]["node1"]["hostname"]),
                    ("Node1 CPU Count", data_source["nodes"]["node1"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_source["nodes"]["node1"] else 0,
                     data_target["nodes"]["node1"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_target["nodes"]["node1"] else 0),
                    ("Node1 CPU Model", data_source["nodes"]["node1"]["hw_info"]["cpu"]["model"] if "hw_info" in data_source["nodes"]["node1"] else 0,
                     data_target["nodes"]["node1"]["hw_info"]["cpu"]["model"] if "hw_info" in data_target["nodes"]["node1"] else "None"),
                    ("Node1 Memory Size (GB)", round(data_source["nodes"]["node1"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_source["nodes"]["node1"] else 0,
                     round(data_target["nodes"]["node1"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_target["nodes"]["node1"] else 0),
                    ("Node1 Interface Count", len(data_source["nodes"]["node1"]["interfaces"]) if "interfaces" in data_source["nodes"]["node1"] else 0,
                     len(data_target["nodes"]["node1"]["interfaces"]) if "interfaces" in data_target["nodes"]["node1"] else 0),
                    ("Node1 OS Name", data_source["nodes"]["node1"]["hw_info"]["os"]["name"] if "hw_info" in data_source["nodes"]["node1"] else "None",
                     data_target["nodes"]["node0"]["hw_info"]["os"]["name"] if "hw_info" in data_target["nodes"]["node0"] else "None"),
                    ("Node1 OS Version", data_source["nodes"]["node1"]["hw_info"]["os"]["version"] if "hw_info" in data_source["nodes"]["node1"] else "None",
                     data_target["nodes"]["node1"]["hw_info"]["os"]["version"] if "hw_info" in data_target["nodes"]["node1"] else "None"),
                ]
            )

            if "hw_info" in data_source["nodes"]["node1"]:
                for storage_source, storage_target in zip(data_source["nodes"]["node1"]["hw_info"]["storage"], data_target["nodes"]["node1"]["hw_info"]["storage"]):
                    source_node1_storage_size = storage_source["size_gb"]
                    target_node1_storage_size = storage_target["size_gb"]
                    table_data_infrastructure.append((f"Node1 Storage {storage_source["name"]} Size (GB)", source_node1_storage_size, target_node1_storage_size))

            table_data_infrastructure.extend(
                [
                    ("Node2 Hostname", data_source["nodes"]["node2"]["hostname"], data_target["nodes"]["node2"]["hostname"]),
                    ("Node2 CPU Count", data_source["nodes"]["node2"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_source["nodes"]["node2"] else 0,
                     data_target["nodes"]["node2"]["hw_info"]["cpu"]["cpus"] if "hw_info" in data_target["nodes"]["node2"] else 0),
                    ("Node2 CPU Model", data_source["nodes"]["node2"]["hw_info"]["cpu"]["model"] if "hw_info" in data_source["nodes"]["node2"] else 0,
                     data_target["nodes"]["node2"]["hw_info"]["cpu"]["model"] if "hw_info" in data_target["nodes"]["node2"] else "None"),
                    ("Node2 Memory Size (GB)", round(data_source["nodes"]["node2"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_source["nodes"]["node2"] else 0,
                     round(data_target["nodes"]["node2"]["hw_info"]["memory"]["size_mb"] / 1024) if "hw_info" in data_target["nodes"]["node2"] else 0),
                    ("Node2 Interface Count", len(data_source["nodes"]["node2"]["interfaces"]) if "interfaces" in data_source["nodes"]["node2"] else 0,
                     len(data_target["nodes"]["node2"]["interfaces"]) if "interfaces" in data_target["nodes"]["node2"] else 0),
                    ("Node2 OS Name", data_source["nodes"]["node2"]["hw_info"]["os"]["name"] if "hw_info" in data_source["nodes"]["node2"] else "None",
                     data_target["nodes"]["node0"]["hw_info"]["os"]["name"] if "hw_info" in data_target["nodes"]["node0"] else "None"),
                    ("Node2 OS Version", data_source["nodes"]["node2"]["hw_info"]["os"]["version"] if "hw_info" in data_source["nodes"]["node2"] else "None",
                     data_target["nodes"]["node2"]["hw_info"]["os"]["version"] if "hw_info" in data_target["nodes"]["node2"] else "None"),
                ]
            )

            if "hw_info" in data_source["nodes"]["node2"]:
                for storage_source, storage_target in zip(data_source["nodes"]["node2"]["hw_info"]["storage"], data_target["nodes"]["node2"]["hw_info"]["storage"]):
                    source_node2_storage_size = storage_source["size_gb"]
                    target_node2_storage_size = storage_target["size_gb"]
                    table_data_infrastructure.append((f"Node2 Storage {storage_source["name"]} Size (GB)", source_node2_storage_size, target_node2_storage_size))

        # Node0 Interface computation
        source_node0_interfaces = list()
        if data_source["kind"] == c.F5XC_SITE_TYPE_SMS_V1:
            for interface in data_source["nodes"]["node0"]["interfaces"]:
                interface_details = dict()
                if "dedicated_interface" in interface.keys():
                    interface_details["is_primary"] = "true" if "is_primary" in interface["dedicated_interface"].keys() else "false"
                    interface_details["device_name"] = interface["dedicated_interface"]["device"]
                    interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                    interface_details["interface_type"] = "dedicated_interface"
                    _interfaces = [f"Node0", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                    source_node0_interfaces.append(_interfaces)
                if "ethernet_interface" in interface.keys():
                    interface_details["mtu"] = interface["ethernet_interface"]["mtu"]
                    interface_details["is_primary"] = True if "is_primary" in interface["ethernet_interface"].keys() else False
                    interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                    interface_details["device_name"] = interface["ethernet_interface"]["device"]
                    interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                    if "dhcp_server" in interface["ethernet_interface"].keys():
                        network_prefixes = list()
                        for network in interface["ethernet_interface"]["dhcp_server"]["dhcp_networks"]:
                            network_prefixes.append(network["network_prefix"])
                        interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                    interface_details["interface_type"] = "ethernet_interface"
                    interface_details["segment_network"] = interface["ethernet_interface"]["segment_network"]["name"] if "segment_network" in interface["ethernet_interface"].keys() else "None"
                    _interface = ["Node0", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                    source_node0_interfaces.append(_interface)
        else:
            # Legacy sites
            for interface_name, interface_attrs in data_source["nodes"]["node0"]["interfaces"].items():
                interface_details = dict()
                interface_details["device_name"] = interface_name
                interface_details["ipv4"] = interface_attrs["subnet_param"]["ipv4"] if "subnet_param" in interface_attrs else None
                interface_details["ipv6"] = interface_attrs["subnet_param"]["ipv6"] if "subnet_param" in interface_attrs else None
                interface_details["existing_subnet_id"] =  interface_attrs["existing_subnet_id"] if "existing_subnet_id" in interface_attrs else None
                _interface = [f"Node0", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                source_node0_interfaces.append(_interface)

        target_node0_interfaces = list()
        if data_target["kind"] == c.F5XC_SITE_TYPE_SMS_V2:
            for interface in data_target["nodes"]["node0"]["interfaces"]:
                interface_details = dict()
                interface_details["mtu"] = interface["mtu"] if "mtu" in interface else "None"
                interface_details["is_primary"] = interface["is_primary"] if "is_primary" in interface else "None"
                interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                interface_details["is_management"] = interface["is_management"] if "is_management" in interface else "None"
                interface_details["site_local_network"] = "True" if "site_local_network" in interface["network_option"] else "None"
                interface_details["site_local_inside_network"] = "True" if "site_local_inside_network" in interface["network_option"] else "None"
                interface_details["dhcp_client"] = "True" if "dhcp_client" in interface else "False"
                interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                interface_details["segment_network"] = interface["network_option"]["segment_network"]["name"] if "segment_network" in interface["network_option"].keys() else "None"
                if "dhcp_server" in interface.keys():
                    network_prefixes = list()
                    for network in interface["dhcp_server"]["dhcp_networks"]:
                        network_prefixes.append(network["network_prefix"])
                    interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                if "ethernet_interface" in interface.keys():
                    interface_details["device_name"] = interface["ethernet_interface"]["device"]
                    interface_details["interface_type"] = "ethernet_interface"
                    interface_details["mac"] = interface["ethernet_interface"]["mac"] if "ethernet_interface" in interface["ethernet_interface"].keys() else "None"
                _interface = [f"{interface["name"]}", join_dict_items(interface_details)]
                target_node0_interfaces.append(_interface)

        table_data_infrastructure_interfaces = []
        if len(source_node0_interfaces) >= len(target_node0_interfaces):
            for source, target in itertools.zip_longest(source_node0_interfaces, target_node0_interfaces, fillvalue=["N/A", "N/A"]):
                table_data_infrastructure_interfaces.append(tuple(source + target))
        else:
            for source, target in itertools.zip_longest(source_node0_interfaces, target_node0_interfaces, fillvalue=["Node0", "N/A", "N/A"]):
                table_data_infrastructure_interfaces.append(tuple(source + target))

        if data_source["main_node_count"] > 1 and data_target["main_node_count"] > 1:
            # Node1 Interface computation
            source_node1_interfaces = list()
            if data_source["kind"] == c.F5XC_SITE_TYPE_SMS_V1:
                for interface in data_source["nodes"]["node1"]["interfaces"]:
                    interface_details = dict()
                    if "dedicated_interface" in interface.keys():
                        interface_details["is_primary"] = "true" if "is_primary" in interface["dedicated_interface"].keys() else "false"
                        interface_details["device_name"] = interface["dedicated_interface"]["device"]
                        interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                        interface_details["interface_type"] = "dedicated_interface"
                        _interfaces = [f"Node1", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                        source_node1_interfaces.append(_interfaces)
                    if "ethernet_interface" in interface.keys():
                        interface_details["mtu"] = interface["ethernet_interface"]["mtu"]
                        interface_details["is_primary"] = True if "is_primary" in interface["ethernet_interface"].keys() else False
                        interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                        interface_details["device_name"] = interface["ethernet_interface"]["device"]
                        interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                        if "dhcp_server" in interface["ethernet_interface"].keys():
                            network_prefixes = list()
                            for network in interface["ethernet_interface"]["dhcp_server"]["dhcp_networks"]:
                                network_prefixes.append(network["network_prefix"])
                            interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                        interface_details["interface_type"] = "ethernet_interface"
                        interface_details["segment_network"] = interface["ethernet_interface"]["segment_network"]["name"] if "segment_network" in interface["ethernet_interface"].keys() else "None"
                        _interface = ["Node1", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                        source_node1_interfaces.append(_interface)
            else:
                # Legacy sites
                for interface_name, interface_attrs in data_source["nodes"]["node1"]["interfaces"].items():
                    interface_details = dict()
                    interface_details["device_name"] = interface_name
                    interface_details["ipv4"] = interface_attrs["subnet_param"]["ipv4"] if "subnet_param" in interface_attrs else None
                    interface_details["ipv6"] = interface_attrs["subnet_param"]["ipv6"] if "subnet_param" in interface_attrs else None
                    interface_details["existing_subnet_id"] = interface_attrs["existing_subnet_id"] if "existing_subnet_id" in interface_attrs else None
                    _interface = [f"Node1", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                    source_node1_interfaces.append(_interface)

            target_node1_interfaces = list()
            if data_target["kind"] == c.F5XC_SITE_TYPE_SMS_V2:
                for interface in data_target["nodes"]["node1"]["interfaces"]:
                    interface_details = dict()
                    interface_details["mtu"] = interface["mtu"] if "mtu" in interface else "None"
                    interface_details["is_primary"] = interface["is_primary"] if "is_primary" in interface else "None"
                    interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                    interface_details["is_management"] = interface["is_management"] if "is_management" in interface else "None"
                    interface_details["site_local_network"] = "True" if "site_local_network" in interface["network_option"] else "None"
                    interface_details["site_local_inside_network"] = "True" if "site_local_inside_network" in interface["network_option"] else "None"
                    interface_details["dhcp_client"] = "True" if "dhcp_client" in interface else "False"
                    interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                    interface_details["segment_network"] = interface["network_option"]["segment_network"]["name"] if "segment_network" in interface[
                        "network_option"].keys() else "None"
                    if "dhcp_server" in interface.keys():
                        network_prefixes = list()
                        for network in interface["dhcp_server"]["dhcp_networks"]:
                            network_prefixes.append(network["network_prefix"])
                        interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                    if "ethernet_interface" in interface.keys():
                        interface_details["device_name"] = interface["ethernet_interface"]["device"]
                        interface_details["interface_type"] = "ethernet_interface"
                        interface_details["mac"] = interface["ethernet_interface"]["mac"] if "ethernet_interface" in interface["ethernet_interface"].keys() else "None"
                    _interface = [f"{interface["name"]}", join_dict_items(interface_details)]
                    target_node1_interfaces.append(_interface)

            if len(source_node1_interfaces) >= len(target_node1_interfaces):
                for source, target in itertools.zip_longest(source_node1_interfaces, target_node1_interfaces, fillvalue=["N/A", "N/A"]):
                    table_data_infrastructure_interfaces.append(tuple(source + target))
            else:
                for source, target in itertools.zip_longest(source_node1_interfaces, target_node1_interfaces, fillvalue=["Node1", "N/A", "N/A"]):
                    table_data_infrastructure_interfaces.append(tuple(source + target))

            # Node2 Interface computation
            source_node2_interfaces = list()
            if data_source["kind"] == c.F5XC_SITE_TYPE_SMS_V1:
                for interface in data_source["nodes"]["node2"]["interfaces"]:
                    interface_details = dict()
                    if "dedicated_interface" in interface.keys():
                        interface_details["is_primary"] = "true" if "is_primary" in interface["dedicated_interface"].keys() else "false"
                        interface_details["device_name"] = interface["dedicated_interface"]["device"]
                        interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                        interface_details["interface_type"] = "dedicated_interface"
                        _interfaces = [f"Node2", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                        source_node2_interfaces.append(_interfaces)
                    if "ethernet_interface" in interface.keys():
                        interface_details["mtu"] = interface["ethernet_interface"]["mtu"]
                        interface_details["is_primary"] = True if "is_primary" in interface["ethernet_interface"].keys() else False
                        interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                        interface_details["device_name"] = interface["ethernet_interface"]["device"]
                        interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                        if "dhcp_server" in interface["ethernet_interface"].keys():
                            network_prefixes = list()
                            for network in interface["ethernet_interface"]["dhcp_server"]["dhcp_networks"]:
                                network_prefixes.append(network["network_prefix"])
                            interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                        interface_details["interface_type"] = "ethernet_interface"
                        interface_details["segment_network"] = interface["ethernet_interface"]["segment_network"]["name"] if "segment_network" in interface["ethernet_interface"].keys() else "None"
                        _interface = ["Node2", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                        source_node2_interfaces.append(_interface)
            else:
                # Legacy sites
                for interface_name, interface_attrs in data_source["nodes"]["node2"]["interfaces"].items():
                    interface_details = dict()
                    interface_details["device_name"] = interface_name
                    interface_details["ipv4"] = interface_attrs["subnet_param"]["ipv4"] if "subnet_param" in interface_attrs else None
                    interface_details["ipv6"] = interface_attrs["subnet_param"]["ipv6"] if "subnet_param" in interface_attrs else None
                    interface_details["existing_subnet_id"] = interface_attrs["existing_subnet_id"] if "existing_subnet_id" in interface_attrs else None
                    _interface = [f"Node2", f"{interface_details["device_name"]}", join_dict_items(interface_details)]
                    source_node2_interfaces.append(_interface)

            target_node2_interfaces = list()
            if data_target["kind"] == c.F5XC_SITE_TYPE_SMS_V2:
                for interface in data_target["nodes"]["node2"]["interfaces"]:
                    interface_details = dict()
                    interface_details["mtu"] = interface["mtu"] if "mtu" in interface else "None"
                    interface_details["is_primary"] = interface["is_primary"] if "is_primary" in interface else "None"
                    interface_details["description"] = interface["description"] if interface["description"] != "" else "None"
                    interface_details["is_management"] = interface["is_management"] if "is_management" in interface else "None"
                    interface_details["site_local_network"] = "True" if "site_local_network" in interface["network_option"] else "None"
                    interface_details["site_local_inside_network"] = "True" if "site_local_inside_network" in interface["network_option"] else "None"
                    interface_details["dhcp_client"] = "True" if "dhcp_client" in interface else "False"
                    interface_details["dhcp_server"] = "true" if "dhcp_server" in interface.keys() else "false"
                    interface_details["segment_network"] = interface["network_option"]["segment_network"]["name"] if "segment_network" in interface[
                        "network_option"].keys() else "None"
                    if "dhcp_server" in interface.keys():
                        network_prefixes = list()
                        for network in interface["dhcp_server"]["dhcp_networks"]:
                            network_prefixes.append(network["network_prefix"])
                        interface_details["dhcp_networks"] = ",".join(network_prefixes) if network_prefixes else "None"
                    if "ethernet_interface" in interface.keys():
                        interface_details["device_name"] = interface["ethernet_interface"]["device"]
                        interface_details["interface_type"] = "ethernet_interface"
                        interface_details["mac"] = interface["ethernet_interface"]["mac"] if "ethernet_interface" in interface["ethernet_interface"].keys() else "None"
                    _interface = [f"{interface["name"]}", join_dict_items(interface_details)]
                    target_node2_interfaces.append(_interface)

            if len(source_node2_interfaces) >= len(target_node2_interfaces):
                for source, target in itertools.zip_longest(source_node2_interfaces, target_node2_interfaces, fillvalue=["N/A", "N/A"]):
                    table_data_infrastructure_interfaces.append(tuple(source + target))
            else:
                for source, target in itertools.zip_longest(source_node2_interfaces, target_node2_interfaces, fillvalue=["Node2", "N/A", "N/A"]):
                    table_data_infrastructure_interfaces.append(tuple(source + target))

        ws_infrastructure.append([f"Infrastructure comparison: {data_source["metadata"]["name"]} with {data_target["metadata"]["name"]}"])
        ws_infrastructure.merge_cells(f"A{ws_infrastructure.max_row}:F{ws_infrastructure.max_row}")
        for cell in ws_infrastructure[ws_infrastructure.max_row]:
            cell.fill = GREY_FILL
            cell.font = HEADER_FONT
            cell.alignment = LEFT_CENTER_ALIGNMENT

        # Nodes section
        ws_infrastructure.append(["Hardware/Software", ""])

        for cell in ws_infrastructure[ws_infrastructure.max_row]:
            cell.fill = GREY_FILL_SECTION
            cell.font = SECTION_FONT

        header = ["Item", "Source", "Target"]
        ws_infrastructure.append(header)

        for cell in ws_infrastructure[ws_infrastructure.max_row]:
            cell.fill = GREY_FILL_SECTION
            if cell.column != 1:
                cell.alignment = LEFT_ALIGNMENT

        append_count = 0
        for item, source, target in table_data_infrastructure:
            ws_infrastructure.append([item, source, target])
            append_count = append_count + 1

        start = 4
        end = ws_infrastructure.max_row + 1

        for row_num in range(start, end):
            value_cell_source = ws_infrastructure[row_num][1]
            value_cell_target = ws_infrastructure[row_num][2]
            value_cell_source.alignment = RIGHT_ALIGNMENT
            value_cell_target.alignment = RIGHT_ALIGNMENT

            if value_cell_source.value != value_cell_target.value:
                value_cell_source.font = Font(color='FF8B0000', bold=False)
                value_cell_target.font = Font(color='FF000080', bold=False)

            # Check if the row number is EVEN
            if row_num % 2 == 0:
                # Apply the grey fill to every cell in the current row
                for cell in ws_infrastructure[row_num]:
                    cell.fill = LIGHT_GREY_FILL

        # Nodes interfaces section
        ws_infrastructure.append(["Interfaces", ""])

        for cell in ws_infrastructure[ws_infrastructure.max_row]:
            cell.fill = GREY_FILL_SECTION
            cell.font = SECTION_FONT

        header = ["Node", "Source Interface", "Values", "Target Interface", "Values"]
        ws_infrastructure.append(header)

        for cell in ws_infrastructure[ws_infrastructure.max_row]:
            cell.fill = GREY_FILL_SECTION
            if cell.column != 1:
                cell.alignment = LEFT_ALIGNMENT

        append_count_interface = 0
        for node, source_iface, source_iface_values, target_iface, target_iface_values in table_data_infrastructure_interfaces:
            ws_infrastructure.append([node, source_iface, source_iface_values, target_iface, target_iface_values])
            append_count_interface = append_count_interface + 1

        start = ws_infrastructure.max_row - append_count_interface + 1
        end = ws_infrastructure.max_row + 1

        for row_num in range(start, end):
            node_cell = ws_infrastructure[row_num][0]
            interface_cell_source = ws_infrastructure[row_num][1]
            interface_cell_target = ws_infrastructure[row_num][3]
            value_cell_source = ws_infrastructure[row_num][2]
            value_cell_target = ws_infrastructure[row_num][4]

            node_cell.alignment = Alignment(horizontal="center", vertical='center')
            interface_cell_source.alignment = Alignment(horizontal="center", vertical='center')
            interface_cell_target.alignment = Alignment(horizontal="center", vertical='center')

            if value_cell_source.value != value_cell_target.value:
                value_cell_source.font = Font(color='FF8B0000', bold=False)
                value_cell_target.font = Font(color='FF000080', bold=False)

            # Check if the row number is EVEN
            if row_num % 2 == 0:
                # Apply the grey fill to every cell in the current row
                for cell in ws_infrastructure[row_num]:
                    cell.fill = LIGHT_GREY_FILL

            max_height_for_row = 0

            for col_letter in ["C", "E"]:
                cell = ws_infrastructure[f'{col_letter}{row_num}']

                if cell.value == "N/A":
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')

                cell_c = ws_infrastructure[f'C{row_num}']
                cell_e = ws_infrastructure[f'E{row_num}']
                required_height = max(len(str(cell_c.value)), len(str(cell_e.value))) + 10

                if required_height is not None and required_height > max_height_for_row:
                    max_height_for_row = required_height

                if max_height_for_row > 0:
                    ws_infrastructure.row_dimensions[row_num].height = max_height_for_row

    def build_compare_service(self, order: int = None, data_source: dict = None, data_target: dict = None):
        """

        Parameters
        ----------
        order
        data_source
        data_target

        Returns
        -------

        """

        # WS Services Comparison Tab
        ws_services = self.wb.create_sheet("Services", order)
        ws_services.column_dimensions['A'].width = 25
        ws_services.column_dimensions['B'].width = 80
        ws_services.column_dimensions['C'].width = 80
        ws_services.column_dimensions['D'].width = 20

        source_ns = list()
        target_ns = list()
        source_lbs = list()
        source_ops = list()
        target_lbs = list()
        target_ops = list()
        source_proxies = list()
        target_proxies = list()

        if "namespaces" in data_source:
            for namespace in data_source["namespaces"]:
                source_ns.append(namespace)
            for source_item in data_source["namespaces"].values():
                if "loadbalancer" in source_item.keys():
                    for source_lb_type in source_item["loadbalancer"].keys():
                        source_lbs.extend(list(source_item["loadbalancer"][source_lb_type].keys()))

                if "proxys" in source_item.keys():
                    for source_proxy_type in source_item["proxys"].keys():
                        source_proxies.append(source_item["proxys"][source_proxy_type]["metadata"]["name"])

            for source_item in data_source["namespaces"].values():
                if "origin_pools" in source_item.keys():
                    source_ops.extend(list(source_item["origin_pools"].keys()))

        if "namespaces" in data_target:
            for namespace in data_target["namespaces"]:
                target_ns.append(namespace)
            for target_item in data_target["namespaces"].values():
                if "loadbalancer" in target_item.keys():
                    for target_lb_type in target_item["loadbalancer"].keys():
                        target_lbs.extend(list(target_item["loadbalancer"][target_lb_type].keys()))

                if "proxys" in target_item.keys():
                    for target_proxy_type in target_item["proxys"].keys():
                        target_proxies.append(target_item["proxys"][target_proxy_type]["metadata"]["name"])

            for target_item in data_target["namespaces"].values():
                if "origin_pools" in target_item.keys():
                    target_ops.extend(list(target_item["origin_pools"].keys()))

        table_data_services = [
            ('NS', "\n".join(source_ns) if len(source_ns) > 0 else "None", "\n".join(target_ns) if len(target_ns) > 0 else "None"),
            ('LB', "\n".join(source_lbs) if len(source_lbs) > 0 else "None", "\n".join(target_lbs) if len(target_lbs) > 0 else "None"),
            ('OP', "\n".join(source_ops) if len(source_ops) > 0 else "None", "\n".join(target_ops) if len(target_ops) > 0 else "None"),
            ('EFP', "\n".join(data_source["efp"].keys()) if "efp" in data_source else "None", "\n".join(data_target["efp"].keys()) if "efp" in data_target else "None"),
            ('FPP', "\n".join(data_source["fpp"].keys()) if "fpp" in data_source else "None", "\n".join(data_target["fpp"].keys()) if "fpp" in data_target else "None"),
            ('SMG', "\n".join(data_source["smg"].keys()) if len(data_source["smg"]) > 0 else "None", "\n".join(data_target["smg"].keys()) if len(data_target["smg"]) > 0 else "None"),
            ('DCCG', "\n".join(data_source["dc_cluster_group"].keys()) if "dc_cluster_group" in data_source else "None", "\n".join(data_target["dc_cluster_group"].keys()) if "dc_cluster_group" in data_target else "None"),
            ('Proxies', "\n".join(source_proxies) if len(source_proxies) > 0 else "None", "\n".join(target_proxies) if len(target_proxies) > 0 else "None"),
            ('Segments', "\n".join(data_source["segments"].keys()) if "segments" in data_source else "None", "\n".join(data_target["segments"].keys()) if "segments" in data_target else "None"),
            ('BGP Policies', "\n".join(data_source["bgp"].keys()) if "bgp" in data_source else "None", "\n".join(data_target["bgp"].keys()) if "bgp" in data_target else "None"),
            ('Virtual Sites', "\n".join(data_source["vsites"]) if len(data_source["vsites"]) > 0 else "None", "\n".join(data_target["vsites"]) if len(data_target["vsites"]) else "None"),
        ]

        ws_services.append([f"Services comparison: {data_source["metadata"]["name"]} with {data_target["metadata"]["name"]}"])
        ws_services.merge_cells(f"A{ws_services.max_row}:F{ws_services.max_row}")
        for cell in ws_services[ws_services.max_row]:
            cell.fill = GREY_FILL
            cell.font = HEADER_FONT
            cell.alignment = LEFT_CENTER_ALIGNMENT

        ws_services.append(["Services", ""])

        for cell in ws_services[ws_services.max_row]:
            cell.fill = GREY_FILL_SECTION
            cell.font = SECTION_FONT

        header = ["Item", "Source", "Target"]
        ws_services.append(header)

        for cell in ws_services[ws_services.max_row]:
            cell.fill = GREY_FILL_SECTION
            if cell.column != 1:
                cell.alignment = LEFT_ALIGNMENT

        append_count = 0
        for item, source_value, target_value in table_data_services:
            ws_services.append([item, source_value, target_value])
            append_count += 1

        for row_num in range(ws_services.max_row - append_count + 1, ws_services.max_row + 1):
            value_cell_source = ws_services[row_num][1]
            value_cell_target = ws_services[row_num][2]
            value_cell_source.alignment = LEFT_ALIGNMENT
            value_cell_target.alignment = LEFT_ALIGNMENT

            if value_cell_source.value != value_cell_target.value:
                value_cell_source.font = Font(color='FF8B0000', bold=False)
                value_cell_target.font = Font(color='FF000080', bold=False)

            # Check if the row number is EVEN
            if row_num % 2 == 0:
                # Apply the grey fill to every cell in the current row
                for cell in ws_services[row_num]:
                    cell.fill = LIGHT_GREY_FILL

            max_height_for_row = 0

            for col_letter in ["B", "C"]:
                cell = ws_services[f'{col_letter}{row_num}']
                cell.alignment = Alignment(wrap_text=True, vertical='top')

                cell_b = ws_services[f'B{row_num}']
                cell_c = ws_services[f'C{row_num}']
                required_height = max(len(str(cell_b.value)), len(str(cell_c.value))) + 10

                if required_height is not None and required_height > max_height_for_row:
                    max_height_for_row = required_height

                if max_height_for_row > 0:
                    ws_services.row_dimensions[row_num].height = max_height_for_row

        for col_letter in ["A"]:
            for row_num in range(4, ws_services.max_row + 1):
                cell = ws_services[f'{col_letter}{row_num}']
                cell.alignment = Alignment(horizontal='center', vertical='center')

    def build_compare(self, data_source: dict = None, data_target: dict = None):
        """

        Parameters
        ----------
        data_source: dict
        data_target: dict

        Returns
        -------

        """

        self.build_compare_summary(0, data_source, data_target)
        self.build_compare_infrastructure(1, data_source, data_target)
        self.build_compare_service(2, data_source, data_target)
